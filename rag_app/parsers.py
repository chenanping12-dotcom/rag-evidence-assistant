from __future__ import annotations

import hashlib
import re
from pathlib import Path

from .models import Page, ParsedDocument


SUPPORTED_SUFFIXES = {".md", ".markdown", ".txt", ".pdf", ".docx"}


def _doc_id(path: Path) -> str:
    return hashlib.sha1(str(path.resolve()).encode("utf-8")).hexdigest()[:12]


def _title_from_text(path: Path, text: str) -> str:
    for line in text.splitlines():
        candidate = re.sub(r"^#+\s*", "", line.strip())
        if candidate:
            return candidate[:120]
    return path.stem


def _parse_text(path: Path) -> ParsedDocument:
    text = path.read_text(encoding="utf-8", errors="replace")
    return ParsedDocument(_doc_id(path), _title_from_text(path, text), str(path), path.suffix, [Page(1, text)])


def _parse_pdf(path: Path) -> ParsedDocument:
    from pypdf import PdfReader

    pages = [Page(i + 1, page.extract_text() or "") for i, page in enumerate(PdfReader(str(path)).pages)]
    return ParsedDocument(_doc_id(path), path.stem, str(path), path.suffix, pages)


def _parse_docx(path: Path) -> ParsedDocument:
    from docx import Document

    text = "\n".join(p.text for p in Document(str(path)).paragraphs if p.text.strip())
    return ParsedDocument(_doc_id(path), _title_from_text(path, text), str(path), path.suffix, [Page(1, text)])


def parse_file(path: str | Path) -> ParsedDocument:
    file_path = Path(path)
    if file_path.suffix.lower() not in SUPPORTED_SUFFIXES:
        raise ValueError(f"Unsupported document type: {file_path.suffix}")
    if file_path.suffix.lower() == ".pdf":
        return _parse_pdf(file_path)
    if file_path.suffix.lower() == ".docx":
        return _parse_docx(file_path)
    return _parse_text(file_path)


def parse_directory(directory: str | Path) -> list[ParsedDocument]:
    root = Path(directory)
    if not root.is_dir():
        raise NotADirectoryError(root)
    documents: list[ParsedDocument] = []
    for path in sorted(root.rglob("*")):
        if path.is_file() and path.suffix.lower() in SUPPORTED_SUFFIXES:
            documents.append(parse_file(path))
    if not documents:
        raise ValueError(f"No supported documents found in {root}")
    return documents
